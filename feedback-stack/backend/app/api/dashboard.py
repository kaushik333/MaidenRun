import io
import json
from datetime import UTC, datetime
from pathlib import Path
from typing import Literal

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import HTMLResponse, Response
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel

from app.core.config import settings
from app.core.logging import logger
from app.db import (
    activate_questionnaire,
    create_questionnaire,
    get_connection,
    list_agent_versions,
    list_feedback,
    list_questionnaires,
    set_agent_version,
)
from app.questions import parse_questionnaire_upload

# Served only by the dashboard container (port 8003, see app/main.py) — never
# by the public form — and every route requires the dashboard token. The
# /api/admin/* routes are how the MaidenRun control plane publishes questionnaire
# versions and announces agent versions without redeploying this stack.
router = APIRouter()

_DASHBOARD_HTML_PATH = Path(__file__).resolve().parent.parent / "templates" / "dashboard.html"


def _require_token(token: str | None) -> None:
    # An unset settings.dashboard_token never matches anything — the
    # dashboard is unreachable rather than silently open if it's ever
    # missing, instead of failing open.
    if not settings.dashboard_token or token != settings.dashboard_token:
        raise HTTPException(status_code=403, detail="Missing or invalid dashboard token")


def _summarize(rows: list[dict]) -> list[dict]:
    """Per-question aggregation, grouped by question *text* rather than a
    fixed schema — each row carries its own question/options, so this works
    for any questionnaire version, or a mix of them."""
    by_question: dict[str, dict] = {}
    for row in rows:
        for answer in row["answers"]:
            entry = by_question.setdefault(
                answer["question"],
                {"question": answer["question"], "options": answer.get("options"), "counts": {}, "text_answers": []},
            )
            if not answer.get("answer"):
                continue
            if answer.get("options"):
                entry["counts"][answer["answer"]] = entry["counts"].get(answer["answer"], 0) + 1
            else:
                entry["text_answers"].append(answer["answer"])
    return list(by_question.values())


def _load_everything() -> tuple[list[dict], list[dict], list[dict]]:
    with get_connection() as conn:
        return list_questionnaires(conn), list_agent_versions(conn), list_feedback(conn)


@router.get("/", response_class=HTMLResponse)
def dashboard_page(token: str | None = Query(default=None)) -> HTMLResponse:
    _require_token(token)
    return HTMLResponse(_DASHBOARD_HTML_PATH.read_text())


@router.get("/api/dashboard/feedback")
def dashboard_feedback(
    token: str | None = Query(default=None),
    questionnaire_version: int | None = Query(default=None, description="omitted = active version, 0 = all"),
    agent_version: int | None = Query(default=None, description="omitted = all agent versions"),
) -> dict:
    _require_token(token)
    questionnaires, agent_versions, rows = _load_everything()
    if questionnaire_version is None:
        questionnaire_version = next(q["version"] for q in questionnaires if q["is_active"])
    in_version = [r for r in rows if questionnaire_version == 0 or r["questionnaire_version"] == questionnaire_version]
    filtered = [r for r in in_version if agent_version is None or r["agent_version_id"] == agent_version]

    # Same questionnaire version, split by agent version — "did my change help?"
    comparison = [
        {"agent_version_id": agent["id"], "responses": len(group), "summary": _summarize(group)}
        for agent in [*agent_versions, {"id": None}]
        if (group := [r for r in in_version if r["agent_version_id"] == agent["id"]])
    ]
    return {
        "questionnaires": questionnaires,
        "agent_versions": agent_versions,
        "questionnaire_version": questionnaire_version,
        "rows": filtered,
        "summary": _summarize(filtered),
        "comparison": comparison,
    }


def _agent_label(agent: dict | None) -> str:
    if agent is None:
        return "unknown"
    message = f" — {agent['commit_message']}" if agent["commit_message"] else ""
    return f"{agent['commit_sha'][:7]}{message}"


def _to_xlsx(questionnaires: list[dict], agent_versions: list[dict], rows: list[dict]) -> bytes:
    """An Overview sheet plus one sheet per questionnaire version — each with
    exactly that version's questions as columns, oldest response first."""
    agents = {a["id"]: a for a in agent_versions}
    bold = Font(bold=True)
    workbook = Workbook()

    overview = workbook.active
    overview.title = "Overview"
    overview.append(["Deployment", settings.deployment_id])
    overview.append(["Exported at (UTC)", datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S")])
    overview.append([])
    overview.append(["Questionnaire versions"])
    overview.append(["Version", "Status", "Created (UTC)", "Last activated (UTC)", "Questions", "Responses"])
    for q in questionnaires:
        overview.append([
            q["version"], "active" if q["is_active"] else "", q["created_at"], q["activated_at"] or "",
            len(q["questions"]), q["responses"],
        ])
    overview.append([])
    overview.append(["Agent versions"])
    overview.append(["Commit", "Message", "Repo", "Live from (UTC)", "Status", "Responses"])
    for a in agent_versions:
        overview.append([
            a["commit_sha"], a["commit_message"], a["repo_url"], a["activated_at"], "live" if a["is_active"] else "",
            a["responses"],
        ])
    for row in overview.iter_rows():
        if row[0].value in ("Questionnaire versions", "Agent versions", "Version", "Commit"):
            for cell in row:
                cell.font = bold

    for q in questionnaires:
        sheet = workbook.create_sheet(f"Questionnaire v{q['version']}")
        version_rows = sorted(
            (r for r in rows if r["questionnaire_version"] == q["version"]), key=lambda r: r["submitted_at"]
        )
        columns = [question["question"] for question in q["questions"]]
        sheet.append(["Submitted at (UTC)", "Agent version", *columns])
        for cell in sheet[1]:
            cell.font = bold
        for r in version_rows:
            answers = {a["question"]: a.get("answer", "") for a in r["answers"]}
            agent = _agent_label(agents.get(r["agent_version_id"]))
            sheet.append([r["submitted_at"], agent, *(answers.get(c, "") for c in columns)])
        sheet.freeze_panes = "A2"
        for column_cells in sheet.columns:
            sheet.column_dimensions[column_cells[0].column_letter].width = min(
                60, max(14, *(len(str(c.value or "")) + 2 for c in column_cells))
            )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@router.get("/api/dashboard/export")
def dashboard_export(
    token: str | None = Query(default=None), format: Literal["xlsx", "json"] = Query(default="xlsx")
) -> Response:
    """Every response across every questionnaire and agent version — an Excel
    workbook for people, a JSON archive (full version definitions included)
    for anything else. Also what the control plane saves before teardown."""
    _require_token(token)
    questionnaires, agent_versions, rows = _load_everything()
    if format == "xlsx":
        body = _to_xlsx(questionnaires, agent_versions, rows)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        body = json.dumps(
            {
                "deployment_id": settings.deployment_id,
                "exported_at": datetime.now(UTC).isoformat(),
                "questionnaires": questionnaires,
                "agent_versions": agent_versions,
                "responses": sorted(rows, key=lambda r: r["submitted_at"]),
            },
            indent=2,
        )
        media_type = "application/json"
    return Response(
        content=body,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="feedback-{settings.deployment_id}.{format}"'},
    )


# ---- admin API, used by the MaidenRun control plane ----


class NewQuestionnaire(BaseModel):
    questionnaire_json: str
    activate: bool = True


class AgentVersion(BaseModel):
    commit_sha: str
    commit_message: str = ""
    repo_url: str = ""


@router.get("/api/admin/questionnaires")
def admin_list_versions(token: str | None = Query(default=None)) -> dict:
    _require_token(token)
    with get_connection() as conn:
        return {"questionnaires": list_questionnaires(conn), "agent_versions": list_agent_versions(conn)}


@router.post("/api/admin/questionnaires", status_code=201)
def admin_publish_questionnaire(request: NewQuestionnaire, token: str | None = Query(default=None)) -> dict:
    _require_token(token)
    try:
        questions = parse_questionnaire_upload(request.questionnaire_json)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    with get_connection() as conn:
        version = create_questionnaire(conn, [q.model_dump() for q in questions], request.activate)
    logger.info("published questionnaire v%s (active=%s)", version, request.activate)
    return {"version": version}


@router.post("/api/admin/questionnaires/{version}/activate")
def admin_activate_questionnaire(version: int, token: str | None = Query(default=None)) -> dict:
    _require_token(token)
    with get_connection() as conn:
        if not activate_questionnaire(conn, version):
            raise HTTPException(status_code=404, detail=f"No questionnaire version {version}")
    logger.info("activated questionnaire v%s", version)
    return {"version": version}


@router.post("/api/admin/agent-version")
def admin_set_agent_version(request: AgentVersion, token: str | None = Query(default=None)) -> dict:
    _require_token(token)
    with get_connection() as conn:
        agent_version_id = set_agent_version(conn, request.commit_sha, request.commit_message, request.repo_url)
    logger.info("live agent version is now %s", request.commit_sha)
    return {"id": agent_version_id}
