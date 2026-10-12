import io
import json

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.dashboard import router as dashboard_router
from app.core.config import settings
from app.db import ensure_schema, get_connection
from app.main import app as form_app  # SERVICE_ROLE defaults to "form" in tests
from app.questions import DEFAULT_QUESTIONS

dashboard_app = FastAPI()
dashboard_app.include_router(dashboard_router)
dashboard = TestClient(dashboard_app)
form = TestClient(form_app)
TOKEN = {"token": settings.dashboard_token}

V2 = {
    "questions": [
        {"question": "Was the report accurate?", "options": ["Yes", "No"]},
        {"question": "What should it do better?", "optional": True},
    ]
}


def _publish(questionnaire: dict, activate: bool = True):
    return dashboard.post(
        "/api/admin/questionnaires",
        params=TOKEN,
        json={"questionnaire_json": json.dumps(questionnaire), "activate": activate},
    )


def _answer(questions: list[dict], choice: str = "Yes") -> list[dict]:
    """`choice` for every multiple-choice question that offers it (else its
    first option, e.g. on a 1-5 scale), "Fine" for free-text ones."""
    def pick(q: dict) -> str:
        if not q.get("options"):
            return "Fine"
        return choice if choice in q["options"] else q["options"][0]

    return [{**q, "answer": pick(q)} for q in questions]


def _submit(version: int | None = None, choice: str = "Yes"):
    current = form.get("/api/questions").json()
    body = {"answers": _answer(current["questions"], choice)}
    if version is not None:
        body["questionnaire_version"] = version
    return form.post("/api/feedback", json=body)


def _rows() -> list[dict]:
    return dashboard.get("/api/dashboard/feedback", params={**TOKEN, "questionnaire_version": 0}).json()["rows"]


def test_bundled_questionnaire_is_version_one():
    body = form.get("/api/questions").json()
    assert body["version"] == 1
    assert [q["question"] for q in body["questions"]] == [q.question for q in DEFAULT_QUESTIONS]


def test_publishing_a_version_changes_the_form_at_the_same_url_without_a_restart():
    assert _publish(V2).json() == {"version": 2}
    body = form.get("/api/questions").json()
    assert body["version"] == 2
    assert [q["question"] for q in body["questions"]] == ["Was the report accurate?", "What should it do better?"]


def test_a_version_can_be_published_without_activating_it():
    assert _publish(V2, activate=False).json() == {"version": 2}
    assert form.get("/api/questions").json()["version"] == 1


def test_activating_an_older_version_rolls_back():
    _publish(V2)
    assert dashboard.post("/api/admin/questionnaires/1/activate", params=TOKEN).status_code == 200
    assert form.get("/api/questions").json()["version"] == 1
    assert dashboard.post("/api/admin/questionnaires/99/activate", params=TOKEN).status_code == 404


def test_a_form_opened_before_a_new_version_still_records_the_version_it_showed():
    stale_questions = form.get("/api/questions").json()["questions"]  # reviewer opens v1...
    _publish(V2)  # ...the owner publishes v2 meanwhile...
    response = form.post(  # ...and the reviewer submits their v1 answers.
        "/api/feedback", json={"answers": _answer(stale_questions), "questionnaire_version": 1}
    )
    assert response.status_code == 201
    assert _rows()[0]["questionnaire_version"] == 1


def test_submissions_without_a_version_go_to_the_active_one():
    _publish(V2)
    assert _submit().status_code == 201
    assert _rows()[0]["questionnaire_version"] == 2


def test_unknown_version_is_rejected():
    assert _submit(version=42).status_code == 422


def test_admin_api_requires_the_token_and_is_not_on_the_public_form():
    assert dashboard.get("/api/admin/questionnaires").status_code == 403
    assert dashboard.post("/api/admin/questionnaires", json={"questionnaire_json": "{}"}).status_code == 403
    assert dashboard.post("/api/admin/agent-version", json={"commit_sha": "abc"}).status_code == 403
    assert form.get("/api/admin/questionnaires", params=TOKEN).status_code in (404, 405)
    assert form.post("/api/admin/questionnaires/1/activate", params=TOKEN).status_code in (404, 405)


def test_invalid_uploads_are_rejected_with_a_reason():
    def upload(text: str):
        return dashboard.post("/api/admin/questionnaires", params=TOKEN, json={"questionnaire_json": text})

    assert "Not valid JSON" in upload("{not json").json()["detail"]
    assert "no questions" in upload('{"questions": []}').json()["detail"]
    duplicate = {"questions": [{"question": "Same?"}, {"question": "Same?"}]}
    assert "repeated" in upload(json.dumps(duplicate)).json()["detail"]
    assert form.get("/api/questions").json()["version"] == 1  # nothing was published


def test_responses_are_stamped_with_the_live_agent_version():
    _submit()  # before any agent version is announced
    first = dashboard.post(
        "/api/admin/agent-version", params=TOKEN, json={"commit_sha": "a" * 40, "commit_message": "initial"}
    ).json()["id"]
    _submit()
    # Re-announcing the live commit (e.g. a repeated notification) changes nothing.
    again = dashboard.post("/api/admin/agent-version", params=TOKEN, json={"commit_sha": "a" * 40}).json()["id"]
    assert again == first
    second = dashboard.post(
        "/api/admin/agent-version", params=TOKEN, json={"commit_sha": "b" * 40, "commit_message": "fix citations"}
    ).json()["id"]
    _submit(choice="No")

    stamped = [r["agent_version_id"] for r in sorted(_rows(), key=lambda r: r["submitted_at"])]
    assert stamped == [None, first, second]


def test_dashboard_filters_and_compares_agent_versions():
    dashboard.post("/api/admin/agent-version", params=TOKEN, json={"commit_sha": "a" * 40})
    _submit(choice="Yes")
    _submit(choice="Yes")
    dashboard.post("/api/admin/agent-version", params=TOKEN, json={"commit_sha": "b" * 40})
    _submit(choice="No")
    _publish(V2)
    _submit()

    def feedback(**params):
        return dashboard.get("/api/dashboard/feedback", params={**TOKEN, **params}).json()

    assert len(feedback()["rows"]) == 1  # default: the active version (v2) only
    assert len(feedback(questionnaire_version=1)["rows"]) == 3
    assert len(feedback(questionnaire_version=0)["rows"]) == 4  # 0 = every version
    assert len(feedback(questionnaire_version=1, agent_version=1)["rows"]) == 2

    comparison = feedback(questionnaire_version=1)["comparison"]
    assert [(c["agent_version_id"], c["responses"]) for c in comparison] == [(1, 2), (2, 1)]
    first_question = DEFAULT_QUESTIONS[0].question
    counts = [next(s for s in c["summary"] if s["question"] == first_question)["counts"] for c in comparison]
    assert counts == [{"Yes": 2}, {"No": 1}]


def test_excel_export_has_an_overview_and_one_sheet_per_version():
    dashboard.post(
        "/api/admin/agent-version", params=TOKEN, json={"commit_sha": "c" * 40, "commit_message": "first cut"}
    )
    _submit()
    _publish(V2)
    _submit(choice="No")

    response = dashboard.get("/api/dashboard/export", params={**TOKEN, "format": "xlsx"})
    assert response.status_code == 200
    assert response.headers["content-disposition"].startswith("attachment;")
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Overview", "Questionnaire v1", "Questionnaire v2"]

    v1 = list(workbook["Questionnaire v1"].values)
    assert v1[0] == ("Submitted at (UTC)", "Agent version", *(q.question for q in DEFAULT_QUESTIONS))
    assert len(v1) == 2 and v1[1][1] == "ccccccc — first cut"
    v2 = list(workbook["Questionnaire v2"].values)
    assert v2[0] == ("Submitted at (UTC)", "Agent version", "Was the report accurate?", "What should it do better?")
    assert v2[1][2:] == ("No", "Fine")

    overview = [row for row in workbook["Overview"].values if row and row[0] is not None]
    assert ("Deployment", settings.deployment_id) == overview[0][:2]
    assert any(row[0] == 2 and row[1] == "active" for row in overview)


def test_json_export_carries_every_version_definition_and_response():
    _submit()
    _publish(V2)
    _submit()
    archive = dashboard.get("/api/dashboard/export", params={**TOKEN, "format": "json"}).json()
    assert [q["version"] for q in archive["questionnaires"]] == [1, 2]
    assert archive["questionnaires"][1]["questions"][0]["question"] == "Was the report accurate?"
    assert [r["questionnaire_version"] for r in archive["responses"]] == [1, 2]


def test_a_pre_versioning_database_is_upgraded_in_place():
    # Recreate the original single-table schema with an existing response.
    with get_connection() as conn:
        conn.execute("DROP TABLE feedback, questionnaires, agent_versions CASCADE")
        conn.execute(
            "CREATE TABLE feedback (id SERIAL PRIMARY KEY, deployment_id TEXT NOT NULL, "
            "answers JSONB NOT NULL, submitted_at TIMESTAMPTZ NOT NULL DEFAULT now())"
        )
        conn.execute(
            """INSERT INTO feedback (deployment_id, answers) VALUES ('old', '[{"question": "Q", "answer": "A"}]')"""
        )
        conn.commit()
        ensure_schema(conn, [q.model_dump() for q in DEFAULT_QUESTIONS])

    assert form.get("/api/questions").json()["version"] == 1
    assert [r["questionnaire_version"] for r in _rows()] == [1]
